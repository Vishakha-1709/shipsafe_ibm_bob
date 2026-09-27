import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    IBM_BLUE = RGBColor(15, 98, 254)      # #0F62FE
    DARK_TEXT = RGBColor(15, 23, 42)      # #0F172A
    MUTED_TEXT = RGBColor(100, 116, 139)  # #64748B
    CARD_BG = RGBColor(241, 245, 249)     # #F1F5F9
    ACCENT_GREEN = RGBColor(16, 185, 129) # #10B981
    ACCENT_RED = RGBColor(239, 68, 68)    # #EF4444

    def add_header(slide, title_text, category_text="IBM BOB 2.0 HACKATHON"):
        # Category / Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = IBM_BLUE

        # Main Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = DARK_TEXT

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide_layout = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(slide_layout)

    # Background accent card
    bg_shape = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.7), Inches(5.1))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = CARD_BG
    bg_shape.line.color.rgb = RGBColor(226, 232, 240)

    # Title & Subtitle inside card
    t_box = s1.shapes.add_textbox(Inches(1.3), Inches(1.8), Inches(10.7), Inches(3.8))
    tf = t_box.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "🛡️ ShipSafe"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = IBM_BLUE

    p1 = tf.add_paragraph()
    p1.text = "AI Release-Readiness & Pre-Flight Security Gatekeeper"
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = DARK_TEXT
    p1.space_before = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "Turn Idea into Impact Faster • Powered by IBM Bob IDE Workflows"
    p2.font.size = Pt(16)
    p2.font.color.rgb = MUTED_TEXT
    p2.space_before = Pt(8)

    p3 = tf.add_paragraph()
    p3.text = "Team: Quantum Quokkas  |  Track: Developer Productivity & Tooling"
    p3.font.size = Pt(14)
    p3.font.color.rgb = IBM_BLUE
    p3.space_before = Pt(24)

    # -------------------------------------------------------------
    # SLIDE 2: The Problem
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(slide_layout)
    add_header(s2, "The Release Bottleneck in Modern Development")

    # 3 Problem Cards
    card_data = [
        ("⏳ Manual & Inconsistent Audits", "Developers spend 45-60 minutes searching for missing tests, dependencies, and entry points before every release."),
        ("🚨 Accidental Credential Exposure", "Hardcoded API keys, database credentials, and active .env files slip into public GitHub repositories undetected."),
        ("❌ Late-Stage Failure in CI/CD", "Missing documentation, broken configs, and zero test coverage are caught late, halting shipping velocity.")
    ]

    for i, (head, desc) in enumerate(card_data):
        x = Inches(0.8 + (i * 4.0))
        shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.7), Inches(4.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = RGBColor(226, 232, 240)

        tb = s2.shapes.add_textbox(x + Inches(0.2), Inches(2.3), Inches(3.3), Inches(3.7))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = DARK_TEXT

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = MUTED_TEXT
        p_desc.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 3: The Solution - ShipSafe
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(slide_layout)
    add_header(s3, "ShipSafe: Instant Automated Release Assurance")

    # Left: Core Pillars
    shape_left = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.6))
    shape_left.fill.solid()
    shape_left.fill.fore_color.rgb = CARD_BG
    shape_left.line.color.rgb = RGBColor(226, 232, 240)

    tb_l = s3.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.0), Inches(4.0))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    
    p = tf_l.paragraphs[0]
    p.text = "🎯 4-Pillar Release Scoring (0-100)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = IBM_BLUE

    pillars = [
        "📖 Documentation (25 pts): README quality & licensing",
        "🧪 Testing (25 pts): Test suite presence & breadth",
        "🔒 Security Hygiene (25 pts): Regex secret interception",
        "⚙️ Maintainability (25 pts): Dependencies & .gitignore"
    ]
    for pil in pillars:
        p_pil = tf_l.add_paragraph()
        p_pil.text = pil
        p_pil.font.size = Pt(13)
        p_pil.font.color.rgb = DARK_TEXT
        p_pil.space_before = Pt(10)

    # Right: Key Features
    shape_right = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.6))
    shape_right.fill.solid()
    shape_right.fill.fore_color.rgb = CARD_BG
    shape_right.line.color.rgb = RGBColor(226, 232, 240)

    tb_r = s3.shapes.add_textbox(Inches(7.1), Inches(2.2), Inches(5.1), Inches(4.0))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "⚡ Key Capabilities"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = IBM_BLUE

    caps = [
        "⚡ Sub-0.05s Audit Time: Instant local scanning",
        "🚨 Secret Hunter: Catches AWS, OpenAI, DB credentials",
        "📋 Prioritized Checklist: Ranked by critical impact",
        "🛡️ Zip Slip Protected: Hardened against malicious zips"
    ]
    for cap in caps:
        p_cap = tf_r.add_paragraph()
        p_cap.text = cap
        p_cap.font.size = Pt(13)
        p_cap.font.color.rgb = DARK_TEXT
        p_cap.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 4: Before vs After IBM Bob
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(slide_layout)
    add_header(s4, "Quantified Developer Impact: Before vs. After Bob")

    # Before Card
    s_before = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.6))
    s_before.fill.solid()
    s_before.fill.fore_color.rgb = RGBColor(254, 242, 242)
    s_before.line.color.rgb = RGBColor(254, 202, 202)

    tb_b = s4.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.0), Inches(4.0))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "❌ Traditional Manual Process"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED

    b_items = [
        "• 45+ minutes spent manual folder searching",
        "• Security checks are inconsistent and missed",
        "• Release checklist written manually from scratch",
        "• Release readiness is subjective guesswork"
    ]
    for it in b_items:
        p_it = tf_b.add_paragraph()
        p_it.text = it
        p_it.font.size = Pt(13)
        p_it.font.color.rgb = DARK_TEXT
        p_it.space_before = Pt(10)

    # After Card
    s_after = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.6))
    s_after.fill.solid()
    s_after.fill.fore_color.rgb = RGBColor(236, 253, 245)
    s_after.line.color.rgb = RGBColor(167, 243, 208)

    tb_a = s4.shapes.add_textbox(Inches(7.1), Inches(2.2), Inches(5.1), Inches(4.0))
    tf_a = tb_a.text_frame
    tf_a.word_wrap = True
    p = tf_a.paragraphs[0]
    p.text = "✅ ShipSafe + IBM Bob Workflow"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    a_items = [
        "• <0.05s automated repository synthesis",
        "• 100% deterministic regex secret interception",
        "• Automated 4-pillar readiness score & PR checklist",
        "• 98% reduction in pre-release preparation time"
    ]
    for it in a_items:
        p_it = tf_a.add_paragraph()
        p_it.text = it
        p_it.font.size = Pt(13)
        p_it.font.color.rgb = DARK_TEXT
        p_it.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 5: How IBM Bob Was Used (Core Evidence)
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(slide_layout)
    add_header(s5, "IBM Bob IDE: Our Core Development Partner")

    bob_steps = [
        ("1. Architecture Planning", "Planned the 4-pillar modular pipeline and deterministic scanner structure."),
        ("2. Implementation", "Authored regex pattern matchers, token masking, and scoring algorithms."),
        ("3. Security Hardening", "Audited ZIP extraction for Zip Slip vulnerabilities and memory safety."),
        ("4. Unit Test Authoring", "Generated comprehensive automated test suite in tests/test_analyzer.py.")
    ]

    for i, (title, desc) in enumerate(bob_steps):
        x = Inches(0.8 + (i * 2.95))
        shape = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(2.8), Inches(4.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = RGBColor(226, 232, 240)

        tb = s5.shapes.add_textbox(x + Inches(0.15), Inches(2.3), Inches(2.5), Inches(3.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = IBM_BLUE

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = DARK_TEXT
        p_desc.space_before = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 6: Summary & Future Roadmap
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(slide_layout)
    add_header(s6, "Summary & Future Roadmap")

    shape_sum = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.7), Inches(4.6))
    shape_sum.fill.solid()
    shape_sum.fill.fore_color.rgb = CARD_BG
    shape_sum.line.color.rgb = RGBColor(226, 232, 240)

    tb_s = s6.shapes.add_textbox(Inches(1.3), Inches(2.3), Inches(10.7), Inches(3.8))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    p.text = "🚀 Turn Idea into Impact Faster with ShipSafe"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = IBM_BLUE

    points = [
        "✅ Shipped MVP: Working Streamlit app, passing tests, and complete bob_sessions evidence.",
        "⏱️ Measurable Value: Reduces 45 minutes of release overhead down to 0.04 seconds.",
        "🔮 Future Roadmap: GitHub Action integration for automated CI/CD release gatekeeping.",
        "🔗 GitHub Repo: github.com/Vishakha-1709/shipsafe_ibm_bob"
    ]
    for pt in points:
        p_pt = tf_s.add_paragraph()
        p_pt.text = pt
        p_pt.font.size = Pt(14)
        p_pt.font.color.rgb = DARK_TEXT
        p_pt.space_before = Pt(12)

    # Save
    out_path = os.path.join(os.path.dirname(__file__), "ShipSafe_Presentation.pptx")
    prs.save(out_path)
    print(f"Created presentation at {out_path}")

if __name__ == "__main__":
    create_presentation()
